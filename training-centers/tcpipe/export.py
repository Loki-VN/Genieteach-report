"""Xuất SQLite → một file Excel: centers, branches, contacts_personal + dữ liệu thô và nhật ký nguồn."""
from __future__ import annotations

import json
import sqlite3
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .classify import CATEGORY_LABELS

HEADER_FILL = PatternFill("solid", fgColor="E8E3F3")


def _flat(v):
    """Danh sách JSON → chuỗi 'a; b' cho dễ đọc trong Excel."""
    if isinstance(v, str) and v.startswith("["):
        try:
            return "; ".join(str(x) for x in json.loads(v))
        except ValueError:
            return v
    if isinstance(v, (list, tuple)):
        return "; ".join(_flat(x) if isinstance(x, (list, tuple)) else str(x) for x in v)
    if isinstance(v, dict):
        return json.dumps(v, ensure_ascii=False)
    return v


def _sheet(wb: Workbook, title: str, columns: list[str], rows: list[list], widths: dict[str, int] | None = None):
    ws = wb.create_sheet(title)
    ws.append(columns)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = HEADER_FILL
        c.alignment = Alignment(vertical="top", wrap_text=True)
    for r in rows:
        ws.append([_flat(v) for v in r])
    ws.freeze_panes = "B2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(columns, 1):
        width = (widths or {}).get(col)
        if width is None:
            longest = max([len(str(col))] + [len(str(_flat(r[i - 1]) or "")) for r in rows[:300]])
            width = min(max(10, longest + 2), 60)
        ws.column_dimensions[get_column_letter(i)].width = width
    return ws


def _query(conn: sqlite3.Connection, sql: str) -> tuple[list[str], list[list]]:
    cur = conn.execute(sql)
    return [d[0] for d in cur.description], [list(r) for r in cur.fetchall()]


def export_excel(conn: sqlite3.Connection, path: Path, notes: list[str]) -> dict[str, int]:
    wb = Workbook()
    wb.remove(wb.active)
    counts = {}

    cols, rows = _query(conn, "SELECT * FROM centers ORDER BY province_new, brand_name")
    _sheet(wb, "centers", cols, rows)
    counts["centers"] = len(rows)

    cols, rows = _query(conn, "SELECT b.center_id, c.brand_name, b.branch_name, b.address_raw, b.address_normalized, "
                              "b.province_new, b.ward_new, b.phone, b.place_id, b.source_url FROM branches b "
                              "JOIN centers c USING(center_id) ORDER BY b.province_new, c.brand_name")
    _sheet(wb, "branches", cols, rows)
    counts["branches"] = len(rows)

    cols, rows = _query(conn, "SELECT p.center_id, c.brand_name, p.person_name, p.title, p.phone, p.email, "
                              "p.source_url, p.collected_at, CASE WHEN p.title LIKE '%Sở GD&ĐT%' "
                              "THEN 'Danh sách Sở GD&ĐT' ELSE 'Website trung tâm' END AS source_type "
                              "FROM contacts_personal p JOIN centers c USING(center_id) ORDER BY c.brand_name")
    _sheet(wb, "contacts_personal", cols, rows)
    counts["contacts_personal"] = len(rows)

    # Dữ liệu thô từ danh sách Sở (kể cả dòng ngoài phạm vi như tư vấn du học)
    sgd_cols = ["province", "old_province", "center_name", "in_scope", "legal_entity_name", "address_raw", "district",
                "director_name", "director_appointment", "phone", "email", "other_contact", "license_number",
                "license_date", "establish_decision", "license_issuer", "scope", "status_text", "as_of", "list_kind", "center_id",
                "source_url", "list_page", "fetched_at"]
    sgd_rows = []
    for r in conn.execute("SELECT * FROM source_records WHERE source='sgd' ORDER BY province, record_id"):
        p = json.loads(r["payload"])
        extra = p.get("extra") or {}
        sgd_rows.append([r["province"], extra.get("old_province"), p["center_name"], p.get("in_scope", True), p.get("legal_entity_name"),
                         p.get("address_raw"), p.get("district"), p.get("director_name"), p.get("director_appointment"),
                         p.get("phone"), p.get("email"), p.get("other_contact"), p.get("license_number"),
                         p.get("license_date"), p.get("establish_decision"), p.get("license_issuer"), p.get("scope"),
                         p.get("status_text"), extra.get("as_of"), extra.get("list_kind"), r["center_id"], r["source_url"],
                         extra.get("list_page"), r["fetched_at"]])
    _sheet(wb, "nguon_so_gddt", sgd_cols, sgd_rows)

    web_cols = ["url", "final_url", "discovered_from", "province_hint", "title", "category", "subjects", "delivery_mode", "phones", "emails",
                "fanpage_urls", "tax_codes", "branch_addresses", "persons", "claims", "founded_year",
                "last_activity_date", "pages", "errors", "center_id", "fetched_at"]
    web_rows = []
    for r in conn.execute("SELECT * FROM source_records WHERE source='website' ORDER BY record_id"):
        p = json.loads(r["payload"])
        web_rows.append([p["url"], p.get("final_url"), p.get("discovered_from"), p.get("province_hint"), p.get("title"), p.get("category"), p.get("subjects"),
                         p.get("delivery_mode"), [x for x, _ in p.get("phones", [])], [x for x, _ in p.get("emails", [])],
                         p.get("fanpage_urls"), p.get("tax_codes"), [a for a, _ in p.get("branch_addresses", [])],
                         [f"{x['title']}: {x['name']}" for x in p.get("persons", [])],
                         [f"{c['kind']}={c['value']} «{c['quote'][:120]}»" for c in p.get("claims", [])][:5],
                         p.get("founded_year"), p.get("last_activity_date"), p.get("pages"), p.get("errors"),
                         r["center_id"], r["fetched_at"]])
    _sheet(wb, "nguon_website", web_cols, web_rows)

    cols, rows = _query(conn, "SELECT province, source, url, outcome, records, detail, checked_at FROM source_log "
                              "ORDER BY source, province")
    _sheet(wb, "nhat_ky_nguon", cols, rows)

    cols, rows = coverage(conn)
    _sheet(wb, "bao_phu_tinh", cols, rows)

    ws = wb.create_sheet("ghi_chu", 0)
    ws.column_dimensions["A"].width = 140
    ws.append([f"Xuất lúc {datetime.now().strftime('%d/%m/%Y %H:%M')} — pipeline training-centers"])
    ws["A1"].font = Font(bold=True)
    for line in notes:
        ws.append([line])
        ws.cell(ws.max_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return counts


CATS = tuple(CATEGORY_LABELS.values())


def coverage(conn: sqlite3.Connection) -> tuple[list[str], list[list]]:
    """Báo cáo độ phủ theo 34 tỉnh: số trung tâm theo loại, số dòng Sở, nguồn đọc được / bị chặn."""
    from .admin_units import provinces

    cols = ["province_code", "province_new", "centers", *CATS, "branches", "contacts_personal", "sgd_rows",
            "sgd_as_of", "sgd_sources_ok", "sgd_sources_failed", "websites_ok", "sgd_failed_detail"]
    rows = []
    unknown = conn.execute("SELECT count(*) FROM centers WHERE province_new IS NULL").fetchone()[0]
    for prov in provinces():
        c = prov.code
        n = conn.execute("SELECT category, count(*) FROM centers WHERE province_new=? GROUP BY category",
                         (prov.name,)).fetchall()
        by_cat = {r[0]: r[1] for r in n}
        sgd_rows = conn.execute("SELECT payload FROM source_records WHERE source='sgd' AND province=?", (c,)).fetchall()
        as_of = sorted({(json.loads(r[0]).get("extra") or {}).get("as_of") or "" for r in sgd_rows} - {""})
        log = conn.execute("SELECT outcome, url, detail FROM source_log WHERE source='sgd' AND province=?",
                           (c,)).fetchall()
        failed = [f"{r[1]} → {r[0]}" for r in log if r[0] != "ok"]
        rows.append([c, prov.name, sum(by_cat.values()), *[by_cat.get(k, 0) for k in CATS],
                     conn.execute("SELECT count(*) FROM branches WHERE province_new=?", (prov.name,)).fetchone()[0],
                     conn.execute("SELECT count(*) FROM contacts_personal p JOIN centers USING(center_id) "
                                  "WHERE province_new=?", (prov.name,)).fetchone()[0],
                     len(sgd_rows), "; ".join(as_of), sum(1 for r in log if r[0] == "ok"), len(failed),
                     conn.execute("SELECT count(*) FROM source_records s JOIN centers c USING(center_id) "
                                  "WHERE s.source='website' AND c.province_new=?", (prov.name,)).fetchone()[0],
                     "; ".join(failed)[:2000]])
    if unknown:
        rows.append(["", "(chưa xác định tỉnh)", unknown] + [""] * (len(cols) - 3))
    return cols, rows
