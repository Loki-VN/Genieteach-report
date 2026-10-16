"""Nguồn 1: danh sách / quyết định cho phép hoạt động do Sở GD&ĐT công bố.

Trang Sở thường là: bài viết + file đính kèm (.xlsx/.pdf/.doc), hoặc bảng HTML.
Module này: tìm file đính kèm và link danh sách → tải → đọc bảng → nhận diện cột → bản ghi.
"""
from __future__ import annotations

import io
import logging
import re
import zipfile
from dataclasses import asdict, dataclass, field
from urllib.parse import urljoin, urlsplit

from selectolax.lexbor import LexborHTMLParser as HTMLParser

from ..http import FetchError, PoliteClient
from ..normalize import clean_text, extract_emails, fold, normalize_phone

log = logging.getLogger(__name__)

ATTACHMENT_EXT = (".xlsx", ".xls", ".pdf", ".doc", ".docx", ".csv")
_LIST_LINK = re.compile(r"danh sach.*(trung tam|ngoai ngu|tin hoc|ky nang|co so giao duc)|"
                        r"(trung tam|ngoai ngu|tin hoc).*(duoc cap phep|duoc phep|cho phep)")

# Từ khóa nhận diện cột (đã bỏ dấu). Thứ tự quan trọng: khóa cụ thể đứng trước.
# Từ bắt đầu bằng "=" phải khớp nguyên ô; còn lại khớp theo cụm từ trong ô.
COLUMNS: list[tuple[str, tuple[str, ...]]] = [
    ("stt", ("=stt", "=tt", "=so tt")),
    ("establish_decision", ("qd cho phep thanh lap", "quyet dinh cho phep thanh lap", "quyet dinh thanh lap",
                            "qd thanh lap")),
    ("other_contact", ("lien he khac",)),
    ("email", ("dia chi email", "email", "e-mail", "thu dien tu")),
    ("scope", ("mon giang day", "mon hoc", "cho phep thanh lap", "noi dung hoat dong", "noi dung dao tao", "linh vuc", "hoat dong dao tao",
               "chuong trinh dao tao")),
    ("legal_entity_name", ("don vi chu quan", "cong ty chu quan", "co quan chu quan", "to chuc thanh lap",
                           "chu dau tu", "don vi thanh lap", "ten doanh nghiep", "phap nhan")),
    ("director_name", ("giam doc", "nguoi dung dau", "nguoi dai dien", "ho ten", "ho va ten")),
    ("phone", ("dien thoai", "sdt", "so dt", "hotline", "=lien he")),
    ("license_date", ("ngay cap", "ngay ky", "ngay quyet dinh", "thoi gian cap phep", "ngay cho phep")),
    ("license_number", ("qd cap phep", "quyet dinh", "so qd", "giay phep", "so hieu", "=qd")),
    ("address_raw", ("dia chi", "dia diem", "tru so", "noi dat")),
    ("center_name", ("ten trung tam", "ten co so", "ten don vi", "=trung tam", "=ten", "ten goi")),
    ("district", ("quan/huyen", "quan, huyen", "=huyen", "=quan", "huyen (tx, tp)", "huyen (thi xa")),
    ("status_text", ("trang thai", "tinh trang", "ghi chu")),
]
_LICENSE_NO = re.compile(r"(\d{1,6})\s*/\s*((?:Q[\u0110D]|GCN|GP|CN)[\w\-/&]*)", re.I)
_DATE = re.compile(r"(\d{1,2})\s*[/.\-]\s*(\d{1,2})\s*[/.\-]\s*((?:19|20)\d{2})")


def _cell_matches(folded: str, word: str) -> bool:
    if word.startswith("="):
        return folded == word[1:]
    return f" {word}" in f" {folded}"


@dataclass
class SgdRecord:
    center_name: str
    source_url: str
    province_code: str
    legal_entity_name: str | None = None
    address_raw: str | None = None
    district: str | None = None
    director_name: str | None = None
    phone: str | None = None
    license_number: str | None = None
    license_date: str | None = None
    license_issuer: str | None = None
    status_text: str | None = None
    scope: str | None = None             # nội dung được phép hoạt động
    establish_decision: str | None = None
    other_contact: str | None = None
    email: str | None = None
    director_appointment: str | None = None   # QĐ công nhận giám đốc (nếu ghi kèm tên)
    in_scope: bool = True                # False: tư vấn du học… (ngoài phạm vi nghiên cứu)
    extra: dict = field(default_factory=dict)


# ---------------------------------------------------------------------------
# Khám phá link
# ---------------------------------------------------------------------------
_ATTACH_HINT = re.compile(r"danh sach|danh-sach|trung tam|trung-tam|ngoai ngu|ngoai-ngu|tin hoc|ky nang|ttnn|tt-nn|"
                          r"cap phep|cho phep|nd125|nd-125")
_DOWNLOAD_QS = re.compile(r"[?&](download|pdf|file|attachment)=", re.I)


def find_links(html: str, base_url: str) -> tuple[list[dict], list[dict]]:
    """→ (file đính kèm, link tới trang danh sách khác).

    File đính kèm gồm link có đuôi .xlsx/.pdf/… và link tải dạng '?download=1&id=N'. Sắp xếp theo
    mức liên quan: thuộc chính bài viết (cùng đường dẫn) hoặc tên có từ khóa danh sách lên trước;
    file ở thanh bên (lịch tiếp công dân…) xuống cuối.
    """
    tree = HTMLParser(html)
    texts: dict[str, str] = {}
    for a in tree.css("a[href]"):
        href = (a.attributes.get("href") or "").strip()
        if not href or href.startswith(("javascript:", "mailto:", "tel:", "#")):
            continue
        url = urljoin(base_url, href)
        text = clean_text(a.text())
        if len(text) > len(texts.get(url, "")):
            texts[url] = text
        else:
            texts.setdefault(url, text)
    page_path = urlsplit(base_url).path
    attachments, lists = [], []
    for url, text in texts.items():
        parts = urlsplit(url)
        ext = re.search(r"\.(xlsx?|pdf|docx?|csv)$", parts.path.lower())
        is_download = bool(_DOWNLOAD_QS.search("?" + parts.query)) and "download" in parts.query.lower()
        if ext or is_download:
            score = 0
            if parts.path == page_path:
                score += 3
            if _ATTACH_HINT.search(fold(f"{text} {parts.path}")):
                score += 2
            if re.search(r"lich tiep|tiep cong dan|lich cong tac", fold(text)):
                score -= 5
            attachments.append({"url": url, "text": text, "ext": ext.group(1) if ext else "file", "score": score})
        elif _LIST_LINK.search(fold(text)):
            lists.append({"url": url, "text": text})
    attachments.sort(key=lambda a: -a["score"])
    return attachments, lists


# ---------------------------------------------------------------------------
# Đọc bảng
# ---------------------------------------------------------------------------
def html_tables(html: str) -> list[list[list[str]]]:
    tables = []
    for t in HTMLParser(html).css("table"):
        rows = [[clean_text(c.text(separator=" ")) for c in tr.css("td, th")] for tr in t.css("tr")]
        rows = [r for r in rows if any(r)]
        if len(rows) >= 2:
            tables.append(rows)
    return tables


def xlsx_tables(content: bytes) -> list[list[list[str]]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    tables = []
    for ws in wb.worksheets:
        rows = [[clean_text("" if v is None else str(v)) for v in row] for row in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(r)]
        if rows:
            tables.append(rows)
    return tables


def xls_tables(content: bytes) -> list[list[list[str]]]:
    """Excel 97-2003 (.xls)."""
    import xlrd

    book = xlrd.open_workbook(file_contents=content)
    tables = []
    for sh in book.sheets():
        rows = [[clean_text(_xls_cell(v)) for v in sh.row_values(i)] for i in range(sh.nrows)]
        rows = [r for r in rows if any(r)]
        if rows:
            tables.append(rows)
    return tables


def _xls_cell(v) -> str:
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return "" if v is None else str(v)


def pdf_tables(content: bytes) -> list[list[list[str]]]:
    import pdfplumber

    content = content[content.find(b"%PDF"):]
    rows: list[list[str]] = []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for page in pdf.pages:
            for table in page.extract_tables():
                rows.extend([clean_text(c or "") for c in r] for r in table if r and any(r))
    return [rows] if rows else []


def pdf_has_text(content: bytes) -> bool:
    """PDF scan (ảnh) không có lớp chữ → cần OCR, ghi nhận để xử lý tay."""
    import pdfplumber

    with pdfplumber.open(io.BytesIO(content[content.find(b"%PDF"):])) as pdf:
        return any((p.extract_text() or "").strip() for p in pdf.pages[:3])


def detect_header(rows: list[list[str]], max_scan: int = 12) -> tuple[int, dict[int, str]] | None:
    """Tìm dòng tiêu đề: dòng có ≥ 2 cột nhận diện được, trong đó có tên trung tâm."""
    for i, row in enumerate(rows[:max_scan]):
        mapping: dict[int, str] = {}
        for j, cell in enumerate(row):
            f = fold(cell)
            if not f or len(f) > 80:
                continue
            for key, words in COLUMNS:
                if key in mapping.values():
                    continue
                if any(_cell_matches(f, w) for w in words):
                    mapping[j] = key
                    break
        if "center_name" in mapping.values() and len(mapping) >= 2:
            return i, mapping
    return None


def parse_license(text: str | None) -> tuple[str | None, str | None]:
    """'QĐ số 1234/QĐ-SGDĐT ngày 12/3/2024' → ('1234/QĐ-SGDĐT', '2024-03-12')."""
    if not text:
        return None, None
    m = _LICENSE_NO.search(re.sub(r"\s*-\s*", "-", text))   # "1170/QĐ- UBND" → "1170/QĐ-UBND"
    number = f"{m.group(1)}/{m.group(2)}" if m else None
    d = _DATE.search(text)
    date = f"{d.group(3)}-{int(d.group(2)):02d}-{int(d.group(1)):02d}" if d else None
    return number, date


_STATUS_IN_CELL = [
    (re.compile(r"dinh chi"), "đình chỉ"),
    (re.compile(r"giai the|thu hoi|cham dut hoat dong"), "giải thể/thu hồi"),
    (re.compile(r"dung hoat dong|ngung hoat dong|tam dung"), "dừng hoạt động"),
]
def _split_director(text: str | None) -> tuple[str | None, str | None]:
    """'Đặng Thị Dư (670/QĐ-SGD ĐT ngày 17/6/2024)' → ('Đặng Thị Dư', '670/QĐ-SGD ĐT ngày 17/6/2024')."""
    if not text:
        return None, None
    m = re.match(r"^\s*([^()]+?)\s*\((.*)\)?\s*$", text)
    if m:
        return clean_text(m.group(1)), clean_text(m.group(2).rstrip(")"))
    return clean_text(text), None


_OUT_OF_SCOPE = re.compile(r"tu van du hoc|van phong dai dien|du hoc")


def rows_to_records(rows: list[list[str]], source_url: str, province_code: str,
                    issuer: str | None = None) -> tuple[list[SgdRecord], dict]:
    found = detect_header(rows)
    if not found:
        return [], {"header": None}
    h, mapping = found
    meta = {"header_row": h, "columns": {rows[h][j]: k for j, k in mapping.items()}}
    records = []
    for row in rows[h + 1:]:
        again = detect_header([row], max_scan=1)       # bảng nhiều phần lặp lại dòng tiêu đề
        if again:
            mapping = again[1]
            continue
        if sum(1 for c in row if c) < 3:               # mảnh chữ do PDF tách dòng
            continue
        values = {key: row[j] for j, key in mapping.items() if j < len(row) and row[j]}
        name = values.get("center_name")
        if not name or fold(name) in ("ten trung tam", "(2)", "2") or re.fullmatch(r"\(?\d+\)?", name):
            continue
        number, date = parse_license(values.get("license_number"))
        number2, date2 = parse_license(values.get("license_date"))
        number = number or number2
        status = values.get("status_text")
        row_text = fold(" ".join(row))
        for pattern, label in _STATUS_IN_CELL:
            if pattern.search(row_text):
                status = f"{label}: {status}" if status else label
                break
        director, appointment = _split_director(values.get("director_name"))
        emails = extract_emails((values.get("email") or "").replace(" ", ""))
        rec = SgdRecord(
            center_name=name, source_url=source_url, province_code=province_code,
            legal_entity_name=values.get("legal_entity_name"), address_raw=values.get("address_raw"),
            district=values.get("district"), director_name=director, director_appointment=appointment,
            email=emails[0] if emails else None,
            phone=normalize_phone(values.get("phone")), license_number=number or values.get("license_number"),
            license_date=date2 or date, license_issuer=issuer, status_text=status, scope=values.get("scope"),
            establish_decision=values.get("establish_decision"), other_contact=values.get("other_contact"),
            in_scope=not _OUT_OF_SCOPE.search(fold(f"{name} {values.get('scope') or ''}")),
        )
        records.append(rec)
    return records, meta


# ---------------------------------------------------------------------------
# Khảo sát (GĐ0): tải trang nguồn, liệt kê định dạng, đọc thử bảng
# ---------------------------------------------------------------------------
def docx_tables(content: bytes) -> list[list[list[str]]]:
    import docx

    document = docx.Document(io.BytesIO(content))
    tables = []
    for t in document.tables:
        rows = [[clean_text(c.text) for c in r.cells] for r in t.rows]
        rows = [r for r in rows if any(r)]
        if len(rows) >= 2:
            tables.append(rows)
    return tables


def sniff_format(content: bytes, content_type: str, url: str) -> str:
    """Nhận dạng định dạng theo chữ ký file (link tải '?download=1' không có đuôi)."""
    head = content[:8]
    if b"%PDF" in content[:1024]:          # có cổng chèn vài byte rác trước chữ ký %PDF
        return "pdf"
    if head.startswith(b"PK\x03\x04"):
        names = zipfile.ZipFile(io.BytesIO(content)).namelist()
        if any(n.startswith("xl/") for n in names):
            return "xlsx"
        if any(n.startswith("word/") for n in names):
            return "docx"
        return "zip"
    if head.startswith(b"Rar!") or head.startswith(b"7z\xbc\xaf"):
        return "rar/7z"         # file nén: cần giải nén thủ công
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        return "ole"            # .xls / .doc đời cũ
    ct = (content_type or "").lower()
    if "html" in ct or content.lstrip()[:15].lower().startswith((b"<!doctype", b"<html")):
        return "html"
    return url.lower().rsplit(".", 1)[-1] if "." in url[-6:] else "unknown"


def tables_from(content: bytes, content_type: str, url: str) -> tuple[str, list[list[list[str]]], list[str]]:
    fmt = sniff_format(content, content_type, url)
    notes = []
    if fmt == "pdf":
        if not pdf_has_text(content):
            notes.append("PDF scan không có lớp chữ — cần OCR/xử lý tay")
            return "pdf_scan", [], notes
        return "pdf", pdf_tables(content), notes
    if fmt == "xlsx":
        return fmt, xlsx_tables(content), notes
    if fmt == "docx":
        return fmt, docx_tables(content), notes
    if fmt == "ole":
        try:
            return "xls", xls_tables(content), notes
        except Exception:      # không phải .xls → .doc đời cũ
            notes.append("File .doc (Word 97-2003) chưa hỗ trợ đọc tự động — cần mở tay")
            return "doc", [], notes
    if fmt == "html":
        return fmt, html_tables(content.decode("utf-8", errors="replace")), notes
    notes.append(f"Định dạng {fmt} chưa hỗ trợ đọc tự động")
    return fmt, [], notes


def probe_source(client: PoliteClient, url: str, province_code: str, issuer: str,
                 follow: int = 10, follow_lists: int = 5) -> dict:
    """Đọc một URL nguồn Sở: chính trang đó, các file đính kèm liên quan; nếu chưa ra bản ghi nào thì đi
    tiếp tối đa `follow_lists` trang danh sách được link tới (một tầng) và file đính kèm của chúng."""
    out: dict = {"url": url, "robots_allowed": None, "pages": [], "records": 0, "samples": [], "errors": [],
                 "_records": []}
    try:
        out["robots_allowed"] = client.allowed(url)
        page = client.get(url)
    except FetchError as e:
        out["errors"].append(f"{type(e).__name__}: {e}")
        return out
    all_records: list[SgdRecord] = []
    seen = {url}

    def fetch_attachments(html: str, base: str, limit: int) -> tuple[list, list[dict]]:
        attachments, lists = find_links(html, base)
        got = []
        for att in [a for a in attachments if a["score"] > 0 and a["url"] not in seen][:limit]:
            seen.add(att["url"])
            try:
                got.append((att["url"], client.get(att["url"])))
            except FetchError as e:
                out["errors"].append(f"{att['url']}: {type(e).__name__}: {e}")
        return got, (attachments, lists)

    def parse(queue: list) -> None:
        for src_url, res in queue:
            fmt, tables, notes = tables_from(res.content, res.content_type, src_url)
            info = {"url": src_url, "format": fmt, "http_status": res.status, "tables": len(tables), "notes": notes}
            for t in tables:
                recs, meta = rows_to_records(t, src_url, province_code, issuer)
                info.setdefault("headers", []).append(meta)
                all_records.extend(recs)
            out["pages"].append(info)

    queue = [(url, page)]
    lists: list[dict] = []
    if "html" in page.content_type:
        got, (attachments, lists) = fetch_attachments(page.text, page.final_url, follow)
        out["attachments"] = attachments
        out["list_links"] = lists[:20]
        queue += got
    parse(queue)
    if not all_records and follow_lists:
        for lst in [x for x in lists if x["url"] not in seen][:follow_lists]:
            seen.add(lst["url"])
            try:
                sub = client.get(lst["url"])
            except FetchError as e:
                out["errors"].append(f"{lst['url']}: {type(e).__name__}: {e}")
                continue
            more = [(lst["url"], sub)]
            if "html" in sub.content_type:
                more += fetch_attachments(sub.text, sub.final_url, follow)[0]
            parse(more)
    out["records"] = len(all_records)
    out["samples"] = [asdict(r) for r in all_records[:5]]
    filled = {}
    for key in ("legal_entity_name", "address_raw", "director_name", "phone", "license_number", "license_date"):
        filled[key] = round(sum(1 for r in all_records if getattr(r, key)) / len(all_records), 2) if all_records else None
    out["fill_rate"] = filled
    out["_records"] = all_records
    return out
